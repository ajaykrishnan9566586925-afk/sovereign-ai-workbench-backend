from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any
import xml.etree.ElementTree as ET

try:
    import fitz  # PyMuPDF
except ImportError:  # optional until installed
    fitz = None

try:
    from openpyxl import load_workbook
except ImportError:  # optional until installed
    load_workbook = None

try:
    import pandas as pd
except ImportError:  # optional legacy Excel support
    pd = None

try:
    from PIL import Image
except ImportError:  # optional until installed
    Image = None

try:
    import pytesseract
except ImportError:  # optional OCR dependency
    pytesseract = None


SUPPORTED_EXTENSIONS = {
    ".pdf", ".txt", ".md", ".markdown", ".csv", ".tsv", ".json",
    ".docx", ".xlsx", ".xlsm", ".pptx", ".png", ".jpg", ".jpeg", ".tif", ".tiff",
}


@dataclass
class ExtractedDocument:
    filename: str
    extension: str
    text: str
    page_count: int
    ocr_used: bool = False


def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _xml_text(blob: bytes) -> str:
    root = ET.fromstring(blob)
    return " ".join(part.strip() for part in root.itertext() if part and part.strip())


def _extract_docx(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        names = [name for name in archive.namelist() if name == "word/document.xml" or name.startswith("word/header") or name.startswith("word/footer")]
        return "\n".join(_xml_text(archive.read(name)) for name in names)


def _extract_pptx(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        names = sorted(name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml"))
        return "\n".join(_xml_text(archive.read(name)) for name in names)


def _extract_spreadsheet(path: Path) -> str:
    if path.suffix.lower() == ".xls":
        if pd is None:
            raise RuntimeError("Legacy .xls support requires pandas and xlrd. Install the project requirements.")
        sheets = pd.read_excel(path, sheet_name=None, header=None)
        sections = []
        for name, frame in sheets.items():
            rows = [" | ".join(row) for row in frame.fillna("").astype(str).values.tolist()]
            sections.append(f"Sheet: {name}\n" + "\n".join(rows))
        return "\n\n".join(sections)
    if load_workbook is None:
        raise RuntimeError("Excel support requires openpyxl. Install the project requirements.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sections: list[str] = []
    for sheet in workbook.worksheets:
        rows: list[str] = [f"Sheet: {sheet.title}"]
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                rows.append(" | ".join(values))
        sections.append("\n".join(rows))
    workbook.close()
    return "\n\n".join(sections)


def _extract_delimited(path: Path) -> str:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as handle:
        rows = csv.reader(handle, delimiter=delimiter)
        return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)


def _extract_image(path: Path) -> tuple[str, bool]:
    if Image is None or pytesseract is None:
        return "Image uploaded. OCR is unavailable because Pillow or Tesseract is not installed.", False
    try:
        return pytesseract.image_to_string(Image.open(path)), True
    except Exception as exc:
        return f"Image uploaded, but OCR failed: {exc}", False


def extract_document(path: Path, filename: str) -> ExtractedDocument:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type {suffix or '(none)'}. Supported types: {allowed}")
    if path.stat().st_size == 0:
        raise ValueError("The uploaded file is empty.")

    ocr_used = False
    page_count = 1
    if suffix == ".pdf":
        if fitz is None:
            raise RuntimeError("PDF support requires PyMuPDF. Install the project requirements.")
        pages: list[str] = []
        with fitz.open(path) as pdf:
            page_count = len(pdf)
            for page in pdf:
                text = page.get_text("text").strip()
                if not text and Image is not None and pytesseract is not None:
                    pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
                    text = pytesseract.image_to_string(image)
                    ocr_used = True
                pages.append(f"[Page {len(pages) + 1}]\n{text}")
        content = "\n\n".join(pages)
    elif suffix in {".docx"}:
        content = _extract_docx(path)
    elif suffix == ".pptx":
        content = _extract_pptx(path)
    elif suffix in {".xlsx", ".xlsm"}:
        content = _extract_spreadsheet(path)
    elif suffix in {".csv", ".tsv"}:
        content = _extract_delimited(path)
    elif suffix == ".json":
        content = json.dumps(json.loads(path.read_text(encoding="utf-8", errors="ignore")), indent=2, ensure_ascii=False)
    elif suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        content, ocr_used = _extract_image(path)
    else:
        content = path.read_text(encoding="utf-8", errors="ignore")

    content = _clean(content)
    if not content:
        raise ValueError(f"No readable text was found in {filename}.")
    return ExtractedDocument(filename, suffix.lstrip("."), content, page_count, ocr_used)


def _tokens(text: str) -> set[str]:
    return {token for token in re.findall(r"[a-zA-Z0-9][a-zA-Z0-9_-]{2,}", text.lower())}


def retrieve(query: str, documents: list[ExtractedDocument], top_k: int = 5) -> list[dict[str, Any]]:
    query_tokens = _tokens(query)
    results: list[dict[str, Any]] = []
    for document in documents:
        paragraphs = [part.strip() for part in re.split(r"\n\s*\n|(?<=\.)\s+(?=[A-Z])", document.text) if part.strip()]
        for index, paragraph in enumerate(paragraphs):
            tokens = _tokens(paragraph)
            overlap = len(query_tokens & tokens)
            phrase_bonus = 2 if query.lower() in paragraph.lower() else 0
            score = (overlap + phrase_bonus) / max(1, len(query_tokens))
            if score > 0 or not query_tokens:
                page_match = re.search(r"\[Page (\d+)\]", paragraph)
                page = int(page_match.group(1)) if page_match else (index + 1 if document.extension == "pdf" else None)
                results.append({
                    "file": document.filename,
                    "detail": f"{document.extension.upper()}" + (f" · page {page}" if page else ""),
                    "content": paragraph[:1200],
                    "score": round(min(1.0, score), 4),
                    "page": page,
                    "ocr_used": document.ocr_used,
                })
    results.sort(key=lambda item: item["score"], reverse=True)
    return results[:top_k]


def extractive_answer(query: str, results: list[dict[str, Any]], documents: list[ExtractedDocument]) -> str:
    if not results:
        return "I could not find relevant content in the uploaded files. Try a more specific question or upload a file containing the requested information."
    lead = results[0]["content"]
    if len(results) == 1:
        return f"Based on the uploaded documents, the most relevant finding is:\n\n{lead}"
    bullets = "\n".join(f"- {item['content']}" for item in results[:4])
    return f"Based on the uploaded documents, these passages best answer your question:\n\n{bullets}"


def document_summary(document: ExtractedDocument) -> dict[str, Any]:
    return {
        "name": document.filename,
        "type": document.extension,
        "pages": document.page_count,
        "characters": len(document.text),
        "ocr_used": document.ocr_used,
    }


def answer_question(query: str, documents: list[ExtractedDocument]) -> tuple[str, list[dict[str, Any]]]:
    results = retrieve(query, documents)
    return extractive_answer(query, results, documents), results


def load_bytes(filename: str, data: bytes, root: Path) -> ExtractedDocument:
    safe_name = Path(filename or "uploaded_file").name
    target = root / safe_name
    target.write_bytes(data)
    return extract_document(target, safe_name)


def as_jsonable(documents: list[ExtractedDocument]) -> list[dict[str, Any]]:
    return [document_summary(document) for document in documents]


def json_bytes(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False).encode("utf-8")


def file_like(data: bytes) -> io.BytesIO:
    return io.BytesIO(data)


__all__ = ["SUPPORTED_EXTENSIONS", "ExtractedDocument", "answer_question", "as_jsonable", "extract_document", "load_bytes"]
